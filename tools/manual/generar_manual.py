"""Convierte el manual en Markdown a Word, con el formato de entrega del proyecto.

Entrada:  docs/Manual_Uso_Plataforma_Modulo1_Facilitadores.md
Salida:   Documentos/Documentos de requerimientos/Manual_Uso_Plataforma_Modulo1_Facilitadores.docx
          (el PDF se exporta despues desde Word)

Soporta lo que usa el manual: titulos h1-h4, parrafos, vinetas, listas numeradas,
tablas, citas (recuadros), reglas horizontales y negrita/cursiva/codigo en linea.

Uso:  python generar_manual.py
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8")
AQUI = Path(__file__).resolve().parent
REPO = AQUI.parents[1]
ORIGEN = REPO / "docs" / "Manual_Uso_Plataforma_Modulo1_Facilitadores.md"
DESTINO = REPO.parent / "Documentos" / "Documentos de requerimientos" / "Manual_Uso_Plataforma_Modulo1_Facilitadores.docx"
AZUL = "0B4582"
GRIS = "EEF2F7"


def sombrear(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def escribir(parrafo, texto):
    """Negrita, cursiva y codigo en linea dentro de un parrafo."""
    for trozo in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            parrafo.add_run(trozo[1:-1]).italic = True
        elif trozo.startswith("`") and trozo.endswith("`"):
            run = parrafo.add_run(trozo[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            parrafo.add_run(trozo)


def fila_de(linea):
    return [c.strip() for c in linea.strip().strip("|").split("|")]


def tabla(doc, filas):
    encabezado, cuerpo = filas[0], filas[2:]
    t = doc.add_table(rows=1, cols=len(encabezado))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, texto in enumerate(encabezado):
        celda = t.rows[0].cells[i]
        celda.text = ""
        escribir(celda.paragraphs[0], texto)
        for run in celda.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        sombrear(celda, AZUL)
    for fila in cuerpo:
        celdas = t.add_row().cells
        for i, texto in enumerate(fila[:len(encabezado)]):
            celdas[i].text = ""
            escribir(celdas[i].paragraphs[0], texto)
    doc.add_paragraph()


def cita(doc, lineas):
    """Recuadro gris de una sola celda para los avisos '>'."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    celda = t.rows[0].cells[0]
    celda.text = ""
    for i, linea in enumerate(lineas):
        p = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
        escribir(p, linea)
    sombrear(celda, GRIS)
    doc.add_paragraph()


def campo(parrafo, instruccion):
    """Inserta un campo de Word (indice o numero de pagina)."""
    inicio, texto, fin = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    texto.set(qn("xml:space"), "preserve")
    texto.text = instruccion
    fin.set(qn("w:fldCharType"), "end")
    run = parrafo.add_run()._r
    for elemento in (inicio, texto, fin):
        run.append(elemento)


def pie_de_pagina(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Manual de uso de la plataforma · Módulo 1 · página ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x66, 0x6F, 0x7B)
    campo(p, "PAGE")
    for run in p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x66, 0x6F, 0x7B)


def indice(doc):
    # Parrafo suelto y no un Heading: asi el propio indice no se lista dentro.
    h = doc.add_paragraph()
    run = h.add_run("Contenido")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(AZUL)
    campo(doc.add_paragraph(), r'TOC \o "1-2" \h \z \u')
    doc.add_page_break()


def convertir(md):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)

    pie_de_pagina(doc)
    lineas = md.split("\n")
    i, primero, encabezado, conindice = 0, True, False, False
    while i < len(lineas):
        linea = lineas[i].rstrip()

        if not linea.strip():
            i += 1
            continue

        # Tabla.
        if linea.startswith("|") and i + 1 < len(lineas) and set(lineas[i + 1].replace("|", "").strip()) <= set("-: "):
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                filas.append(fila_de(lineas[i]))
                i += 1
            tabla(doc, filas)
            continue

        # Cita.
        if linea.startswith(">"):
            bloque = []
            while i < len(lineas) and lineas[i].startswith(">"):
                bloque.append(lineas[i].lstrip("> ").rstrip())
                i += 1
            cita(doc, bloque)
            continue

        # Regla horizontal: separa capitulos con salto de pagina. La primera
        # cierra la portada, y despues de ella va el indice.
        if linea.strip() == "---":
            doc.add_page_break()
            if not conindice:
                indice(doc)
                conindice = True
            i += 1
            continue

        # Titulos.
        m = re.match(r"(#{1,4}) (.+)", linea)
        if m:
            nivel = len(m.group(1))
            if primero:
                # Portada: el titulo baja al centro de la hoja y la linea
                # siguiente (subtitulo) se centra tambien, en convertir().
                for _ in range(8):
                    doc.add_paragraph()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(m.group(2))
                run.bold = True
                run.font.size = Pt(26)
                run.font.color.rgb = RGBColor.from_string(AZUL)
                primero = False
                encabezado = True
            else:
                h = doc.add_heading(level=min(nivel, 4))
                escribir(h, m.group(2))
                for run in h.runs:
                    run.font.color.rgb = RGBColor.from_string(AZUL)
            i += 1
            continue

        # Listas.
        m = re.match(r"[-*] (.+)", linea)
        if m:
            escribir(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue
        m = re.match(r"\d+\. (.+)", linea)
        if m:
            escribir(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        if encabezado:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            encabezado = False
        escribir(p, linea.strip())
        i += 1

    return doc


def main():
    doc = convertir(ORIGEN.read_text(encoding="utf-8"))
    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DESTINO)
    print("OK:", DESTINO)


if __name__ == "__main__":
    main()

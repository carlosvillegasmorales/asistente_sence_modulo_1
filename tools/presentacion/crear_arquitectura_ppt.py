import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Colors
    NAVY = RGBColor(24, 43, 73)
    DARK_GRAY = RGBColor(60, 64, 67)
    LIGHT_GRAY = RGBColor(245, 247, 250)
    BORDER_GRAY = RGBColor(220, 224, 230)
    WHITE = RGBColor(255, 255, 255)
    
    # ----------------------------------------------------
    # SLIDE 1: PORTADA
    # ----------------------------------------------------
    slide1 = prs.slides.add_slide(blank_slide_layout)
    
    # Background
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.color.rgb = NAVY
    
    # Title Box
    title_box = slide1.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(11), Inches(3.0))
    tf1 = title_box.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "Arquitectura de la Solución"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Calibri"
    
    p2 = tf1.add_paragraph()
    p2.text = "Plataforma Moodle + Asistencia Inteligente & Cumplimiento SENCE"
    p2.font.size = Pt(22)
    p2.font.color.rgb = RGBColor(180, 205, 237)
    p2.font.name = "Calibri"
    p2.space_before = Pt(14)
    
    p3 = tf1.add_paragraph()
    p3.text = "Visión Macro y Componentes Principales"
    p3.font.size = Pt(15)
    p3.font.color.rgb = RGBColor(200, 210, 225)
    p3.font.name = "Calibri"
    p3.space_before = Pt(24)

    # ----------------------------------------------------
    # SLIDE 2: ARQUITECTURA MACRO (LA SLIDE PRINCIPAL)
    # ----------------------------------------------------
    slide2 = prs.slides.add_slide(blank_slide_layout)
    
    # Light clean background
    bg2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = LIGHT_GRAY
    bg2.line.color.rgb = LIGHT_GRAY

    # Header section
    header_box = slide2.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.3))
    tf_h = header_box.text_frame
    tf_h.word_wrap = True
    
    ph1 = tf_h.paragraphs[0]
    ph1.text = "Arquitectura General del Sistema"
    ph1.font.size = Pt(30)
    ph1.font.bold = True
    ph1.font.color.rgb = NAVY
    ph1.font.name = "Calibri"
    
    ph2 = tf_h.add_paragraph()
    ph2.text = "Flujo macro: Cómo se integran el usuario, Moodle y la Inteligencia Artificial en 4 capas clave"
    ph2.font.size = Pt(14)
    ph2.font.color.rgb = DARK_GRAY
    ph2.font.name = "Calibri"
    ph2.space_before = Pt(6)

    # 4 Cards Configuration
    cards_data = [
        {
            "step": "PASO 1",
            "title": "USUARIOS",
            "subtitle": "Estudiantes y Profesores",
            "tag_bg": RGBColor(41, 128, 185),
            "accent": RGBColor(41, 128, 185),
            "bg": RGBColor(255, 255, 255),
            "points": [
                "Interactúan en la interfaz de Moodle como siempre.",
                "Alumnos: resuelven dudas 24/7 y suben documentos SENCE.",
                "Profesores: supervisan analíticas y alertas tempranas."
            ]
        },
        {
            "step": "PASO 2",
            "title": "MOODLE LMS",
            "subtitle": "El Sistema Central",
            "tag_bg": RGBColor(22, 160, 133),
            "accent": RGBColor(22, 160, 133),
            "bg": RGBColor(255, 255, 255),
            "points": [
                "Aporta el contexto académico real del usuario.",
                "Calificaciones con escala chilena (1.0 - 7.0).",
                "Almacenamiento seguro de archivos y control de accesos."
            ]
        },
        {
            "step": "PASO 3",
            "title": "MOTOR INTELIGENTE",
            "subtitle": "Plugin Especializado",
            "tag_bg": RGBColor(230, 126, 34),
            "accent": RGBColor(230, 126, 34),
            "bg": RGBColor(255, 255, 255),
            "points": [
                "Asistente Pedagógico con 3 perfiles formativos.",
                "Validador SENCE: juzga si los PDFs cumplen requisitos.",
                "Detector de alerta temprana ante riesgo de abandono."
            ]
        },
        {
            "step": "PASO 4",
            "title": "NUBE DE IA",
            "subtitle": "Modelos Flexibles (LLMs)",
            "tag_bg": RGBColor(142, 68, 173),
            "accent": RGBColor(142, 68, 173),
            "bg": RGBColor(255, 255, 255),
            "points": [
                "Google Gemini, Claude, OpenAI o Azure AI.",
                "Arquitectura desacoplada (sin ataduras a una marca).",
                "Ahorro de costos mediante Prompt Caching."
            ]
        }
    ]

    left_margin = Inches(0.8)
    card_width = Inches(2.65)
    card_height = Inches(4.7)
    card_gap = Inches(0.38)
    card_top = Inches(2.1)

    for i, c in enumerate(cards_data):
        c_left = left_margin + i * (card_width + card_gap)
        
        # Outer Card Container
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, card_top, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = c["bg"]
        card.line.color.rgb = BORDER_GRAY
        card.line.width = Pt(1.5)
        
        # Step Tag
        tag = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left + Inches(0.2), card_top + Inches(0.2), Inches(1.0), Inches(0.35))
        tag.fill.solid()
        tag.fill.fore_color.rgb = c["tag_bg"]
        tag.line.fill.background()
        tf_tag = tag.text_frame
        tf_tag.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = c["step"]
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = WHITE
        p_tag.alignment = PP_ALIGN.CENTER

        # Card Text Box
        tbox = slide2.shapes.add_textbox(c_left + Inches(0.18), card_top + Inches(0.65), card_width - Inches(0.36), card_height - Inches(0.8))
        tf_card = tbox.text_frame
        tf_card.word_wrap = True

        # Title
        p_title = tf_card.paragraphs[0]
        p_title.text = c["title"]
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY
        p_title.font.name = "Calibri"

        # Subtitle
        p_sub = tf_card.add_paragraph()
        p_sub.text = c["subtitle"]
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = c["accent"]
        p_sub.font.name = "Calibri"
        p_sub.space_after = Pt(14)

        # Bullets
        for pt in c["points"]:
            p_pt = tf_card.add_paragraph()
            p_pt.text = "• " + pt
            p_pt.font.size = Pt(11.5)
            p_pt.font.color.rgb = DARK_GRAY
            p_pt.font.name = "Calibri"
            p_pt.space_after = Pt(8)

        # Arrow connector between cards
        if i < 3:
            arrow_left = c_left + card_width + Inches(0.08)
            arrow = slide2.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, 
                arrow_left, 
                card_top + Inches(2.1), 
                Inches(0.22), 
                Inches(0.35)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor(180, 190, 205)
            arrow.line.fill.background()

    # Footer note on Slide 2
    footer_box = slide2.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.4))
    tf_f = footer_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Principio de diseño: Human-in-the-loop (la IA propone, el facilitador o tutor siempre tiene el control final)."
    p_f.font.size = Pt(11)
    p_f.font.italic = True
    p_f.font.color.rgb = RGBColor(120, 130, 145)

    # Save
    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "docs",
                            "Arquitectura_Macro_Solucion.pptx")
    prs.save(out_path)
    print(f"Presentation saved successfully to {out_path}")

if __name__ == "__main__":
    create_presentation()
